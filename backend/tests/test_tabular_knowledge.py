"""Tabular knowledge (M3): CSV/XLSX parsing and the structured query engine.

These run on SQLite. Tables are stored with the same ingestion helper the
pipeline uses, so the JSON-path queries are exercised on a second dialect
(PostgreSQL is covered by tests/knowledge/test_access_control.py).
"""

from __future__ import annotations

import io
import uuid
from datetime import date

import pytest
from openpyxl import Workbook

from app.core.config import settings
from app.core.exceptions import NotFoundError, ValidationError
from app.knowledge import extraction, tables, tabular
from app.knowledge.access import Principal
from app.knowledge.errors import DocumentExtractionError
from app.knowledge.ingestion import _store_tables
from app.models.enums import DocumentStatus, KnowledgeBaseStatus
from app.models.knowledge import KnowledgeBase, KnowledgeDocument
from tests.conftest import register_org

_ORDERS = (
    "Order,Customer,City,Amount,Status\n"
    '1,Ada,Lagos,"1,000",paid\n'
    "2,Bola,Kano,2500,paid\n"
    "3,Chi,Kano,4000,pending\n"
    "4,Dayo,Abuja,,paid\n"
    "5,Efe,kano,9500.5,paid\n"
)


# ---------------------------------------------------------------- parsing


def test_csv_types_thousands_and_blanks():
    [table] = tabular.parse_csv(_ORDERS.encode())
    assert table.columns == [
        {"name": "Order", "type": "number"},
        {"name": "Customer", "type": "text"},
        {"name": "City", "type": "text"},
        {"name": "Amount", "type": "number"},
        {"name": "Status", "type": "text"},
    ]
    assert table.rows[0]["Amount"] == 1000
    assert table.rows[3]["Amount"] is None
    assert table.rows[4]["Amount"] == 9500.5
    assert isinstance(table.rows[0]["Order"], int)


def test_csv_headers_delimiter_bom_and_non_finite_values():
    data = "﻿name;;name;score\nx;1;2;nan\ny;3;4;inf\n".encode()
    [table] = tabular.parse_csv(data)
    assert [c["name"] for c in table.columns] == ["name", "column_2", "name_2", "score"]
    # "nan"/"inf" are not numbers: the column stays text and JSON-safe.
    assert {c["name"]: c["type"] for c in table.columns}["score"] == "text"
    assert table.rows[0]["score"] == "nan"


def test_empty_csv_has_no_tables_and_extraction_rejects_it():
    assert tabular.parse_csv(b"\n\n") == []
    with pytest.raises(DocumentExtractionError):
        extraction.extract(b"\n", filename="empty.csv")


def test_table_row_limit_marks_truncation(monkeypatch):
    monkeypatch.setattr(settings, "KNOWLEDGE_MAX_TABLE_ROWS", 2)
    table = tabular.build_table("t", [["a"], [1], [2], [3]])
    assert table is not None
    assert table.truncated is True
    assert [r["a"] for r in table.rows] == [1, 2]


def _xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(["Item", "Qty", "Restocked"])
    ws.append(["Rice", 40, date(2026, 9, 1)])
    ws.append(["Beans", 12, None])
    wb.create_sheet("Empty")
    notes = wb.create_sheet("Notes")
    notes.append(["Note"])
    notes.append(["Prices exclude VAT"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_sheets_become_tables_and_text():
    doc = extraction.extract(_xlsx(), filename="stock.xlsx")
    assert [t.name for t in doc.tables] == ["Stock", "Notes"]  # empty sheet skipped
    stock = doc.tables[0]
    assert {c["name"]: c["type"] for c in stock.columns} == {
        "Item": "text",
        "Qty": "number",
        "Restocked": "text",
    }
    assert stock.rows[0]["Restocked"].startswith("2026-09-01")
    assert "Table: Stock" in doc.text
    assert "Item: Rice; Qty: 40" in doc.text


def test_corrupt_xlsx_is_an_extraction_error():
    with pytest.raises(DocumentExtractionError):
        extraction.extract(b"not a zip", filename="broken.xlsx")


# ---------------------------------------------------------------- querying


async def _seed(client, session_factory) -> dict:
    auth = await register_org(client, "tab@acme.example.com", "Acme")
    org_id = uuid.UUID(auth["organization_id"])
    user_id = uuid.UUID(auth["user"]["id"])
    async with session_factory() as s:
        kb = KnowledgeBase(organization_id=org_id, name="Sales", slug="sales", created_by=user_id)
        s.add(kb)
        await s.flush()
        doc = KnowledgeDocument(
            organization_id=org_id,
            knowledge_base_id=kb.id,
            name="orders.csv",
            checksum="0" * 64,
            status=DocumentStatus.READY.value,
            created_by=user_id,
        )
        s.add(doc)
        await s.flush()
        await _store_tables(s, org_id, doc, kb, tabular.parse_csv(_ORDERS.encode()))
        await s.commit()
    owner = Principal(organization_id=org_id, user_id=user_id, role="OWNER")
    async with session_factory() as s:
        [listed] = await tables.list_tables(s, owner)
    return {"org": org_id, "user": user_id, "kb": kb, "doc": doc, "table": listed, "owner": owner}


async def _q(session_factory, principal, table_id, **spec):
    async with session_factory() as s:
        return await tables.query_table(
            s, principal, uuid.UUID(table_id), tables.TableQuery(**spec)
        )


async def test_filters_aggregates_and_group_by(client, session_factory):
    ctx = await _seed(client, session_factory)
    tid, owner = ctx["table"]["table_id"], ctx["owner"]
    assert ctx["table"]["row_count"] == 5

    paid = await _q(
        session_factory,
        owner,
        tid,
        filters=[{"column": "Status", "op": "eq", "value": "paid"}],
        aggregate={"op": "sum", "column": "Amount"},
    )
    assert paid["rows"] == [{"sum_Amount": 13000.5}]

    kano = await _q(
        session_factory,
        owner,
        tid,
        filters=[{"column": "City", "op": "contains", "value": "KANO"}],
        aggregate={"op": "count"},
    )
    assert kano["rows"] == [{"count": 3}]  # case-insensitive contains

    by_city = await _q(
        session_factory, owner, tid, aggregate={"op": "count"}, group_by="City", limit=1
    )
    assert by_city["rows"] == [{"City": "Kano", "count": 2}]
    assert by_city["truncated"] is True

    big = await _q(
        session_factory,
        owner,
        tid,
        columns=["Customer", "Amount"],
        filters=[{"column": "Amount", "op": "gte", "value": "2500"}],
        order_by={"column": "Amount", "descending": True},
        limit=2,
    )
    assert big["rows"] == [
        {"Customer": "Efe", "Amount": 9500.5},
        {"Customer": "Chi", "Amount": 4000},
    ]
    assert big["matched_rows"] == 3
    assert big["truncated"] is True

    blanks = await _q(session_factory, owner, tid, filters=[{"column": "Amount", "op": "is_null"}])
    assert [r["Customer"] for r in blanks["rows"]] == ["Dayo"]


@pytest.mark.parametrize(
    "spec",
    [
        {"filters": [{"column": "Nope", "op": "eq", "value": "x"}]},
        {"filters": [{"column": "City", "op": "gt", "value": "A"}]},
        {"filters": [{"column": "Amount", "op": "contains", "value": "1"}]},
        {"filters": [{"column": "Amount", "op": "eq", "value": "lots"}]},
        {"filters": [{"column": "City", "op": "eq"}]},
        {"aggregate": {"op": "sum", "column": "City"}},
        {"aggregate": {"op": "avg"}},
        {"group_by": "City"},
    ],
)
async def test_invalid_queries_are_rejected(client, session_factory, spec):
    ctx = await _seed(client, session_factory)
    with pytest.raises(ValidationError):
        await _q(session_factory, ctx["owner"], ctx["table"]["table_id"], **spec)


def test_query_spec_forbids_unknown_fields():
    with pytest.raises(ValueError):
        tables.TableQuery.model_validate({"sql": "DROP TABLE x"})


async def test_tables_follow_access_status_and_scope(client, session_factory):
    ctx = await _seed(client, session_factory)
    tid = uuid.UUID(ctx["table"]["table_id"])
    member = Principal(organization_id=ctx["org"], user_id=uuid.uuid4(), role="MEMBER")
    outsider = Principal(organization_id=uuid.uuid4(), user_id=ctx["user"], role="OWNER")
    non_member = Principal(organization_id=ctx["org"], user_id=uuid.uuid4(), role=None)

    async with session_factory() as s:
        assert len(await tables.list_tables(s, member)) == 1
        assert await tables.list_tables(s, outsider) == []
        assert await tables.list_tables(s, non_member) == []
        # An agent scoped to other knowledge bases cannot reach this table.
        with pytest.raises(NotFoundError):
            await tables.get_table(s, ctx["owner"], tid, knowledge_base_ids=[uuid.uuid4()])

        doc = await s.get(KnowledgeDocument, ctx["doc"].id)
        doc.visibility = "RESTRICTED"
        await s.commit()
        assert await tables.list_tables(s, member) == []
        assert len(await tables.list_tables(s, ctx["owner"])) == 1  # creator keeps access

        kb = await s.get(KnowledgeBase, ctx["kb"].id)
        kb.status = KnowledgeBaseStatus.ARCHIVED.value
        await s.commit()
        with pytest.raises(NotFoundError):
            await tables.get_table(s, ctx["owner"], tid)
