"""Tabular extraction: CSV / XLSX → typed tables (+ text for semantic search).

Structured data is kept as data. Each sheet becomes a table whose columns are
typed ("number" when every non-empty value parses as a number, else "text"), so
filters and aggregates can be answered exactly by the structured query engine
(`app.knowledge.tables`) rather than approximated by vector similarity. A text
rendering is still produced so the same rows are findable by semantic search.

Files are parsed strictly as data: formulas are read as their cached values,
macros are never executed, and sizes are bounded.
"""

from __future__ import annotations

import csv
import io
import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any

from app.core.config import settings
from app.knowledge.errors import DocumentExtractionError

_THOUSANDS = re.compile(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$")


@dataclass
class ExtractedTable:
    name: str
    columns: list[dict[str, str]]  # [{"name": ..., "type": "number" | "text"}]
    rows: list[dict[str, Any]]
    truncated: bool = False
    metadata: dict[str, Any] = field(default_factory=dict)


def _column_names(header: list[Any]) -> list[str]:
    names: list[str] = []
    for i, raw in enumerate(header):
        name = str(raw).strip() if raw not in (None, "") else ""
        name = (name or f"column_{i + 1}")[:64]
        base, n = name, 2
        while name in names:
            name = f"{base}_{n}"[:64]
            n += 1
        names.append(name)
    return names


def _to_number(value: Any) -> float | int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return value if math.isfinite(value) else None
    if not isinstance(value, str):
        return None
    text = value.strip()
    if _THOUSANDS.match(text):
        text = text.replace(",", "")
    try:
        number = float(text)
    except ValueError:
        return None
    if not math.isfinite(number):  # "nan"/"inf" are text, not numbers (and not JSON)
        return None
    looks_integral = "." not in text and "e" not in text.lower()
    return int(number) if number.is_integer() and looks_integral else number


def _normalize_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, bool):
        return str(value).lower()
    return value


def build_table(name: str, raw_rows: list[list[Any]]) -> ExtractedTable | None:
    """Header row + data rows → a typed table. Returns None for an empty sheet."""
    rows = [r for r in raw_rows if any(_normalize_cell(c) is not None for c in r)]
    if not rows:
        return None
    header, body = rows[0], rows[1:]
    width = min(max(len(r) for r in rows), settings.KNOWLEDGE_MAX_TABLE_COLUMNS)
    columns = _column_names(list(header[:width]) + [None] * (width - len(header[:width])))

    truncated = len(body) > settings.KNOWLEDGE_MAX_TABLE_ROWS
    body = body[: settings.KNOWLEDGE_MAX_TABLE_ROWS]
    cells = [[_normalize_cell(r[i]) if i < len(r) else None for i in range(width)] for r in body]

    types: list[str] = []
    for i in range(width):
        values = [row[i] for row in cells if row[i] is not None]
        numeric = bool(values) and all(_to_number(v) is not None for v in values)
        types.append("number" if numeric else "text")

    records: list[dict[str, Any]] = []
    for row in cells:
        record: dict[str, Any] = {}
        for i, col in enumerate(columns):
            value = row[i]
            if types[i] == "number":
                record[col] = _to_number(value) if value is not None else None
            else:
                record[col] = None if value is None else str(value)
        records.append(record)
    return ExtractedTable(
        name=name[:128] or "Sheet1",
        columns=[{"name": c, "type": t} for c, t in zip(columns, types, strict=True)],
        rows=records,
        truncated=truncated,
    )


def parse_csv(data: bytes, name: str = "Sheet1") -> list[ExtractedTable]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    sample = text[:4096]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = list(csv.reader(io.StringIO(text), dialect))
    except csv.Error as exc:
        raise DocumentExtractionError(f"Could not read CSV: {exc}") from exc
    table = build_table(name, rows)
    return [table] if table else []


def parse_xlsx(data: bytes) -> list[ExtractedTable]:
    try:
        from openpyxl import load_workbook

        # data_only: formulas yield cached values; read_only: streaming, bounded.
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise DocumentExtractionError(f"Could not read XLSX: {type(exc).__name__}") from exc
    tables: list[ExtractedTable] = []
    try:
        for sheet in workbook.worksheets:
            limit = settings.KNOWLEDGE_MAX_TABLE_ROWS + 2  # header + overflow detection
            raw: list[list[Any]] = []
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i >= limit:
                    break
                raw.append(list(row))
            table = build_table(sheet.title, raw)
            if table is not None:
                tables.append(table)
    finally:
        workbook.close()
    return tables


def render_text(tables: list[ExtractedTable]) -> str:
    """Readable text for embeddings: one line per row, 'column: value; ...'."""
    parts: list[str] = []
    for table in tables:
        names = [c["name"] for c in table.columns]
        lines = [f"Table: {table.name}", "Columns: " + ", ".join(names)]
        for row in table.rows:
            cells = [f"{k}: {v}" for k, v in row.items() if v is not None]
            if cells:
                lines.append("; ".join(cells))
        parts.append("\n".join(lines))
    return "\n\n".join(parts)
